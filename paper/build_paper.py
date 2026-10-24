# -*- coding: utf-8 -*-
"""Render content.py to IEEE-style .docx and IEEEtran .tex, plus figures and the audit CSV.

Run:  python build_paper.py      (needs python-docx and matplotlib)
Reads ../backend/target/brsr-retrieval-eval.txt (written by BrsrFullRetrievalEvaluationTest).
"""
import csv
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import content as C

HERE = Path(__file__).parent
EVAL = HERE.parent / "backend" / "target" / "brsr-retrieval-eval.txt"

# ---------------------------------------------------------------- audit data
# Manual grading of the top-ranked passage per requirement (R relevant, P partial, N not relevant).
# Anything not listed is R. Single rater: the authors must re-grade before submission.
PARTIAL = {"A-07", "A-17", "B-01", "B-02", "B-04", "B-10", "P5-E09", "P5-L03", "P6-E09", "P9-E03"}
NOT = {"A-05", "A-06", "A-08", "A-11", "B-03", "B-05", "B-12", "P5-E06", "P5-E10", "P5-L04", "P5-L05"}


def load_audit():
    rows, code, title, score = [], None, None, None
    pattern = re.compile(r"^(\S+)\s+score=([\d.]+)\s+pages=(\S+)\s*\|\s*(.*)$")
    lines = EVAL.read_text(encoding="utf-8").splitlines()
    for line in lines:
        m = pattern.match(line)
        if m:
            code, score, pages, title = m.group(1), float(m.group(2)), m.group(3), m.group(4)
            rows.append({"code": code, "title": title, "score": score, "pages": pages.strip(","), "top": ""})
        elif line.strip().startswith("top:") and rows:
            rows[-1]["top"] = line.strip()[4:].strip()
    for r in rows:
        r["grade"] = "N" if r["code"] in NOT else "P" if r["code"] in PARTIAL else "R"
        r["group"] = r["code"].split("-")[0]
    return rows


AUDIT = load_audit()
assert len(AUDIT) == 140, len(AUDIT)
COUNT = lambda rows, g: sum(1 for r in rows if r["grade"] == g)
assert (COUNT(AUDIT, "R"), COUNT(AUDIT, "P"), COUNT(AUDIT, "N")) == (119, 10, 11)
GROUPS = ["A", "B"] + [f"P{i}" for i in range(1, 10)]
BY_GROUP = {g: [r for r in AUDIT if r["group"] == g] for g in GROUPS}
PRINCIPLES = [r for r in AUDIT if r["group"].startswith("P")]


def pct(n, d):
    return f"{100 * n / d:.1f}"


# ---------------------------------------------------------------- tables
REQ_ROWS = [
    ("A (General)", 24, "-", 24, 2), ("B (Management)", 12, "-", 12, 1),
    ("P1", 7, 2, 9, 1), ("P2", 4, 5, 9, 1), ("P3", 15, 6, 21, 2), ("P4", 2, 3, 5, 1),
    ("P5", 10, 5, 15, 2), ("P6", 12, 9, 21, 2), ("P7", 2, 1, 3, 1), ("P8", 4, 6, 10, 1),
    ("P9", 6, 5, 11, 1),
]
assert sum(r[3] for r in REQ_ROWS) == 140 and sum(r[4] for r in REQ_ROWS) == 15

TABLES = {
    "reqset": {
        "caption": "BRSR Requirement Set and Batch Plan (L = 12)",
        "header": ["Section", "Gen./Ess.", "Lead.", "Total", "Calls"],
        "rows": [[a, str(b), str(c), str(d), str(e)] for a, b, c, d, e in REQ_ROWS]
                + [["Total", "98", "42", "140", "15"]],
        "widths": [1.1, 0.7, 0.55, 0.55, 0.55],
    },
    "calls": {
        "caption": "Model Calls per Full Analysis (140 Requirements)",
        "header": ["Strategy", "Calls", "Fits 20/day?"],
        "rows": [["One call per requirement", "140", "No"],
                 ["Section batches, L = 12", "15", "Yes (5 spare)"]],
        "widths": [1.9, 0.6, 0.95],
    },
    "audit": {
        "caption": "Top-1 Retrieval Audit on the Infosys BRSR FY 2024-25",
        "header": ["Group", "n", "Rel.", "Part.", "Not", "P@1"],
        "rows": [
            ["Section A", "24"] + [str(COUNT(BY_GROUP["A"], g)) for g in "RPN"] + [pct(COUNT(BY_GROUP["A"], "R"), 24) + "%"],
            ["Section B", "12"] + [str(COUNT(BY_GROUP["B"], g)) for g in "RPN"] + [pct(COUNT(BY_GROUP["B"], "R"), 12) + "%"],
            ["Principles 1-9", "104"] + [str(COUNT(PRINCIPLES, g)) for g in "RPN"] + [pct(COUNT(PRINCIPLES, "R"), 104) + "%"],
            ["All", "140"] + [str(COUNT(AUDIT, g)) for g in "RPN"] + [pct(COUNT(AUDIT, "R"), 140) + "%"],
        ],
        "widths": [1.1, 0.4, 0.45, 0.45, 0.4, 0.65],
        "note": "P@1: share of requirements whose top passage is relevant (strict).",
    },
}

FIGS = {
    "pipeline": {"file": "fig1_pipeline.png", "caption": "ESGenius analysis pipeline. Dashed boxes are deterministic fail-safe paths."},
    "audit": {"file": "fig2_audit.png", "caption": "Top-1 retrieval audit by BRSR section (n = 140)."},
}


# ---------------------------------------------------------------- figures
def make_figures():
    plt.rcParams["font.family"] = "serif"
    plt.rcParams["font.serif"] = ["Times New Roman", "DejaVu Serif"]

    # Fig. 1 pipeline (vertical, single column)
    fig, ax = plt.subplots(figsize=(3.4, 4.9), dpi=300)
    ax.set_xlim(0, 10); ax.set_ylim(0, 20); ax.axis("off")

    def box(x, y, w, h, text, dashed=False, fill="#f2f2f2"):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.25",
                                    fc=fill, ec="black", lw=0.8, ls="--" if dashed else "-"))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=6.3, linespacing=1.15)

    def arrow(x0, y0, x1, y1):
        ax.annotate("", xy=(x1, y1), xytext=(x0, y0),
                    arrowprops=dict(arrowstyle="-|>", lw=0.8, color="black", shrinkA=0, shrinkB=0))

    box(1.5, 18.2, 7, 1.6, "PDF upload (React UI)\nstored; status PROCESSING")
    box(1.5, 15.7, 7, 1.7, "Per-page text extraction\n(PDFBox, no OCR) \u2192 READY")
    box(1.5, 13.2, 7, 1.7, "Page-bounded chunking\n1,200\u20131,800 chars, 200 overlap")
    box(1.5, 10.5, 7, 1.9, "Lexical retrieval per requirement\nBM25 + aliases + anchor gate\ntop-3 passages with page numbers")
    for y0, y1 in [(18.2, 17.4), (15.7, 14.9), (13.2, 12.4)]:
        arrow(5, y0, 5, y1)
    ax.text(5.2, 9.85, "evidence found?", fontsize=6.3, ha="left", va="center", style="italic")
    arrow(5, 10.5, 5, 9.4)
    arrow(5, 9.4, 1.9, 8.5)
    arrow(5, 9.4, 7.9, 8.5)
    ax.text(1.3, 9.0, "no", fontsize=6.3, ha="center")
    ax.text(8.7, 9.0, "yes", fontsize=6.3, ha="center")
    box(0.2, 6.4, 3.9, 2.1, "Deterministic\nNOT_COVERED\n(no model call)", dashed=True, fill="white")
    box(4.9, 6.4, 4.9, 2.1, "Group by BRSR section;\nbatches of \u2264 12\n(15 calls per 140)")
    arrow(7.35, 6.4, 7.35, 5.6)
    box(4.9, 3.6, 4.9, 2.0, "Gemini classification\nJSON array per batch")
    arrow(7.35, 3.6, 7.35, 2.9)
    box(4.9, 0.9, 4.9, 2.0, "On error / omission:\nHUMAN_REVIEW_\nREQUIRED", dashed=True, fill="white")
    arrow(2.15, 6.4, 2.15, 1.0)
    ax.plot([2.15, 4.9], [1.0, 1.0], lw=0.8, color="black")
    ax.text(2.4, 0.35, "assessments saved \u2192 progress UI \u2192 reviewer \u2192 gap-assessment PDF", fontsize=5.8,
            ha="left", va="center")
    fig.savefig(HERE / FIGS["pipeline"]["file"], bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)

    # Fig. 2 stacked bars per section
    fig, ax = plt.subplots(figsize=(3.4, 2.1), dpi=300)
    rel = [COUNT(BY_GROUP[g], "R") for g in GROUPS]
    par = [COUNT(BY_GROUP[g], "P") for g in GROUPS]
    non = [COUNT(BY_GROUP[g], "N") for g in GROUPS]
    x = range(len(GROUPS))
    ax.bar(x, rel, color="#555555", edgecolor="black", lw=0.4, label="Relevant")
    ax.bar(x, par, bottom=rel, color="#bbbbbb", edgecolor="black", lw=0.4, label="Partial")
    ax.bar(x, non, bottom=[a + b for a, b in zip(rel, par)], color="white", edgecolor="black",
           hatch="////", lw=0.4, label="Not relevant")
    ax.set_xticks(list(x)); ax.set_xticklabels(GROUPS, fontsize=6.5)
    ax.set_ylabel("Requirements", fontsize=7); ax.tick_params(axis="y", labelsize=6.5)
    ax.legend(fontsize=6, frameon=False, loc="upper right")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.savefig(HERE / FIGS["audit"]["file"], bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)


# ---------------------------------------------------------------- inline parsing
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def split_inline(text):
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            out.append(("b", part[2:-2]))
        elif part.startswith("*"):
            out.append(("i", part[1:-1]))
        elif part.startswith("`"):
            out.append(("c", part[1:-1]))
        else:
            out.append(("n", part))
    return out


# ---------------------------------------------------------------- docx
def set_font(run, size=None, bold=None, italic=None, small_caps=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if small_caps is not None:
        run.font.small_caps = small_caps


def add_runs(par, text, size=10, base_italic=False, base_bold=False):
    for kind, chunk in split_inline(text):
        pieces = [chunk] if kind == "c" else re.split(r"(?<=[A-Za-z])(_[a-z0-9]+)", chunk)
        for piece in pieces:
            if not piece:
                continue
            sub = kind != "c" and piece.startswith("_") and len(piece) > 1
            r = par.add_run(piece[1:] if sub else piece)
            set_font(r, size=size, bold=base_bold or kind == "b", italic=base_italic != (kind == "i"),
                     name="Courier New" if kind == "c" else "Times New Roman")
            if sub:
                r.font.subscript = True


def add_sub_runs(par, text, size=10):
    """Equation text with _{sub} and ^{sup} markup."""
    for part in re.split(r"(_\{[^}]*\}|\^\{[^}]*\})", text):
        if not part:
            continue
        if part.startswith("_{"):
            r = par.add_run(part[2:-1]); r.font.subscript = True
        elif part.startswith("^{"):
            r = par.add_run(part[2:-1]); r.font.superscript = True
        else:
            r = par.add_run(part)
        set_font(r, size=size, italic=True)


def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, first=None, keep_next=False):
    pf = par.paragraph_format
    par.alignment = align
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    if first is not None:
        pf.first_line_indent = Inches(first)
    pf.keep_with_next = keep_next


def set_cols(section, num, space_twips=360):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:cols")):
        sectPr.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    sectPr.append(cols)


def cell_borders(table):
    tbl = table._tbl
    props = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "000000")
        borders.append(el)
    props.append(borders)


def roman(n):
    vals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    s = ""
    for v, r in vals:
        while n >= v:
            s += r; n -= v
    return s


def build_docx(path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    s0 = doc.sections[0]
    s0.page_width, s0.page_height = Inches(8.5), Inches(11)
    s0.left_margin = s0.right_margin = Inches(0.625)
    s0.top_margin, s0.bottom_margin = Inches(0.75), Inches(1.0)

    # title block (single column)
    p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 6)
    set_font(p.add_run(C.TITLE), size=22, bold=False)
    n = len(C.AUTHORS)
    t = doc.add_table(rows=3, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, name in enumerate(C.AUTHORS):
        for row, (txt, it, sz) in enumerate([(name, False, 11), (" / ".join(C.AFFILIATION[:1]), True, 9),
                                              (", ".join(C.AFFILIATION[1:]), True, 9)]):
            cell = t.cell(row, i)
            cp = cell.paragraphs[0]; fmt(cp, WD_ALIGN_PARAGRAPH.CENTER)
            set_font(cp.add_run(txt), size=sz, italic=it)
    p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 4, 8)
    set_font(p.add_run(C.MENTOR_LINE), size=9, italic=True)

    # two-column body
    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    set_cols(s1, 2)

    p = doc.add_paragraph(); fmt(p, before=0, after=4)
    set_font(p.add_run("Abstract\u2014"), size=9, bold=True, italic=True)
    add_runs(p, C.ABSTRACT, size=9, base_bold=True)
    p = doc.add_paragraph(); fmt(p, before=0, after=8)
    set_font(p.add_run("Index Terms\u2014"), size=9, bold=True, italic=True)
    set_font(p.add_run(C.KEYWORDS), size=9, bold=True)

    h1n = h2n = 0
    tcount = fcount = 0
    emitted_tables, emitted_figs = set(), set()

    def table_block(key):
        nonlocal tcount
        if key in emitted_tables:
            return
        emitted_tables.add(key); tcount += 1
        spec = TABLES[key]
        cp = doc.add_paragraph(); fmt(cp, WD_ALIGN_PARAGRAPH.CENTER, 6, 0, keep_next=True)
        set_font(cp.add_run(f"TABLE {roman(tcount)}"), size=8, small_caps=False)
        cp = doc.add_paragraph(); fmt(cp, WD_ALIGN_PARAGRAPH.CENTER, 0, 3, keep_next=True)
        set_font(cp.add_run(spec["caption"].upper()), size=8)
        tb = doc.add_table(rows=1 + len(spec["rows"]), cols=len(spec["header"]))
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        tb.autofit = False
        cell_borders(tb)
        for ci, h in enumerate(spec["header"]):
            c = tb.cell(0, ci); c.width = Inches(spec["widths"][ci])
            pp = c.paragraphs[0]; fmt(pp, WD_ALIGN_PARAGRAPH.CENTER)
            set_font(pp.add_run(h), size=8, bold=True)
        for ri, row in enumerate(spec["rows"], start=1):
            last = ri == len(spec["rows"]) and spec["rows"][-1][0] in ("Total", "All")
            for ci, val in enumerate(row):
                c = tb.cell(ri, ci); c.width = Inches(spec["widths"][ci])
                pp = c.paragraphs[0]
                fmt(pp, WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER)
                set_font(pp.add_run(val), size=8, bold=last)
        if spec.get("note"):
            np_ = doc.add_paragraph(); fmt(np_, WD_ALIGN_PARAGRAPH.LEFT, 2, 0)
            set_font(np_.add_run(spec["note"]), size=7.5, italic=True)
        sp = doc.add_paragraph(); fmt(sp, after=4)

    def fig_block(key):
        nonlocal fcount
        if key in emitted_figs:
            return
        emitted_figs.add(key); fcount += 1
        spec = FIGS[key]
        pp = doc.add_paragraph(); fmt(pp, WD_ALIGN_PARAGRAPH.CENTER, 4, 2, keep_next=True)
        pp.add_run().add_picture(str(HERE / spec["file"]), width=Inches(3.35 if key == "audit" else 3.0))
        cp = doc.add_paragraph(); fmt(cp, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 8)
        set_font(cp.add_run(f"Fig. {fcount}. {spec['caption']}"), size=8)

    for blk in C.BLOCKS:
        kind = blk[0]
        if kind == "h1":
            h1n += 1; h2n = 0
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 8, 4, keep_next=True)
            set_font(p.add_run(f"{roman(h1n)}. {blk[1]}".upper()), size=10, small_caps=False)
        elif kind == "h2":
            h2n += 1
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 4, 2, keep_next=True)
            set_font(p.add_run(f"{chr(64 + h2n)}. {blk[1]}"), size=10, italic=True)
        elif kind == "p":
            p = doc.add_paragraph(); fmt(p, first=0.17)
            add_runs(p, blk[1])
        elif kind == "eq":
            p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 3, 3)
            add_sub_runs(p, blk[1], size=9.5)
            set_font(p.add_run(f"    ({blk[3]})"), size=10)
        elif kind == "table":
            table_block(blk[1])
        elif kind == "fig":
            fig_block(blk[1])

    p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 8, 4, keep_next=True)
    set_font(p.add_run("ACKNOWLEDGMENT"), size=10)
    p = doc.add_paragraph(); fmt(p, first=0.17); add_runs(p, C.ACK)

    p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 8, 4, keep_next=True)
    set_font(p.add_run("REFERENCES"), size=10)
    for i, ref in enumerate(C.REFERENCES, start=1):
        p = doc.add_paragraph(); fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 2)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        set_font(p.add_run(f"[{i}]\t"), size=8)
        add_runs(p, ref, size=8)
    doc.save(path)


# ---------------------------------------------------------------- latex
LATEX_ESC = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
             "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
UNI = {"\u2013": "--", "\u2014": "---", "\u00d7": r"$\times$", "\u2265": r"$\geq$", "\u2264": r"$\leq$",
       "\u2192": r"$\rightarrow$", "\u2019": "'", "\u2308": r"$\lceil$", "\u2309": r"$\rceil$",
       "\u230a": r"$\lfloor$", "\u230b": r"$\rfloor$", "\u0304": "", "\u00e9": r"\'e"}


def tex_text(s):
    s = s.replace("L̄", "@@BARL@@")
    s = "".join(LATEX_ESC.get(ch, ch) for ch in s)
    for k, v in UNI.items():
        s = s.replace(k, v)
    s = s.replace("@@BARL@@", r"$\bar{L}$")
    return s


def tex_inline(text):
    out = []
    for kind, chunk in split_inline(text):
        if kind == "b":
            out.append(r"\textbf{" + tex_text(chunk) + "}")
        elif kind == "i":
            out.append(r"\textit{" + tex_text(chunk) + "}")
        elif kind == "c":
            out.append(r"\texttt{" + tex_text(chunk) + "}")
        else:
            out.append(tex_text(chunk))
    s = "".join(out)
    s = s.replace("@@BARL@@", r"$\bar{L}$")
    return s


def build_tex(path):
    L = []
    L.append(r"\documentclass[conference]{IEEEtran}")
    L.append(r"\usepackage[T1]{fontenc}\usepackage[utf8]{inputenc}")
    L.append(r"\usepackage{amsmath,graphicx,booktabs,cite}")
    L.append(r"\begin{document}")
    L.append(r"\title{" + tex_text(C.TITLE) + "}")
    authors = r" \and ".join(
        r"\IEEEauthorblockN{" + a + r"}\IEEEauthorblockA{" + tex_text(C.AFFILIATION[0]) + r"\\" +
        tex_text(", ".join(C.AFFILIATION[1:])) + "}" for a in C.AUTHORS)
    L.append(r"\author{" + authors + "}")
    L.append(r"\maketitle")
    L.append(r"\begin{abstract}" + tex_inline(C.ABSTRACT) + r"\end{abstract}")
    L.append(r"\begin{IEEEkeywords}" + tex_text(C.KEYWORDS) + r"\end{IEEEkeywords}")

    emitted = set()
    for blk in C.BLOCKS:
        kind = blk[0]
        if kind == "h1":
            L.append(r"\section{" + tex_text(blk[1]) + "}")
        elif kind == "h2":
            L.append(r"\subsection{" + tex_text(blk[1]) + "}")
        elif kind == "p":
            L.append(tex_inline(blk[1]) + "\n")
        elif kind == "eq":
            L.append(r"\begin{equation}" + blk[2] + r"\end{equation}")
        elif kind == "table" and blk[1] not in emitted:
            emitted.add(blk[1]); spec = TABLES[blk[1]]
            cols = "l" + "c" * (len(spec["header"]) - 1)
            L.append(r"\begin{table}[!t]\centering\caption{" + tex_text(spec["caption"]) + "}")
            L.append(r"\footnotesize\begin{tabular}{" + cols + r"}\toprule")
            L.append(" & ".join(tex_text(h) for h in spec["header"]) + r" \\ \midrule")
            for row in spec["rows"]:
                L.append(" & ".join(tex_text(v) for v in row) + r" \\")
            L.append(r"\bottomrule\end{tabular}")
            if spec.get("note"):
                L.append(r"\par\vspace{2pt}{\scriptsize\textit{" + tex_text(spec["note"]) + "}}")
            L.append(r"\end{table}")
        elif kind == "fig" and blk[1] not in emitted:
            emitted.add(blk[1]); spec = FIGS[blk[1]]
            L.append(r"\begin{figure}[!t]\centering\includegraphics[width=0.95\columnwidth]{" +
                     spec["file"] + "}\n" + r"\caption{" + tex_text(spec["caption"]) + r"}\end{figure}")
    L.append(r"\section*{Acknowledgment}" + tex_inline(C.ACK))
    L.append(r"\begin{thebibliography}{99}")
    for ref in C.REFERENCES:
        L.append(r"\bibitem{} " + tex_inline(ref))
    L.append(r"\end{thebibliography}")
    L.append(r"\end{document}")
    tex = "\n".join(L)
    # number the bibitems so [n] citations resolve via cite keys b1..bN
    counter = iter(range(1, len(C.REFERENCES) + 1))
    tex = re.sub(r"\\bibitem\{\}", lambda m: r"\bibitem{b%d}" % next(counter), tex)
    # in-text citations "[1]" / "[3], [4]" -> \cite
    tex = re.sub(r"\[(\d+)\]", lambda m: r"\cite{b%s}" % m.group(1), tex)
    Path(path).write_text(tex, encoding="utf-8")


def write_csv(path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["requirement_code", "requirement_title", "top_score", "pages", "top_passage_excerpt",
                    "grade_by_assistant (R/P/N)", "grade_by_author_1", "grade_by_author_2"])
        for r in AUDIT:
            w.writerow([r["code"], r["title"], r["score"], r["pages"], r["top"], r["grade"], "", ""])


if __name__ == "__main__":
    make_figures()
    build_docx(HERE / "ESGenius_IEEE_Paper.docx")
    build_tex(HERE / "ESGenius_IEEE_Paper.tex")
    write_csv(HERE / "retrieval_audit.csv")
    print("ok")
